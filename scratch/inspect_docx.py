import sys
import io
import docx

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

doc = docx.Document(r"Bao_Cao/BaoCao_DoAn_BigData_Nhom6_IEEE.docx")
print("=== TỔNG QUAN FILE DOCX ===")
print("Số đoạn (Paragraphs):", len(doc.paragraphs))
print("Số bảng (Tables):", len(doc.tables))
print("\n=== CÁC ĐOẠN TỪ 20 ĐẾN 55 ===")
for i in range(20, min(55, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    txt = p.text.strip()
    if txt:
        print(f"[{i:02d}] {txt}")

print("\n=== TOÀN BỘ CÁC BẢNG HIỆN CÓ ===")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- BẢNG {t_idx+1} ({len(table.rows)} dòng x {len(table.columns)} cột) ---")
    for r_idx, row in enumerate(table.rows):
        row_txt = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
        print(f"  Dòng {r_idx}: {' | '.join(row_txt)}")
